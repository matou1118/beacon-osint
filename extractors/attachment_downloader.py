"""附件下载与文本提取管线（服务器友好版）。

核心原则：下载到临时目录 → 提取文本 → 立即删除二进制文件 → 文本存DB。
不长期占用磁盘空间。需要原文时用 source_url 重新下载。

用法（在 main.py 中自动调用）：
    from extractors.attachment_downloader import process_attachments
    attachment_result = process_attachments(records)
"""
from __future__ import annotations

import hashlib
import re
import tempfile
from pathlib import Path
from typing import Any

try:
    import httpx
except ImportError:
    httpx = None

try:
    # Import the pymupdf name, not the deprecated `fitz` alias: importing `fitz`
    # prints a deprecation warning to STDOUT, which corrupts the JSON that
    # main.py writes to stdout and breaks every consumer that parses it.
    import pymupdf as fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx  # python-docx
except ImportError:
    docx = None

try:
    import openpyxl  # Excel
except ImportError:
    openpyxl = None


MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB 上限
TIMEOUT_SECONDS = 30
MAX_TEXT_LENGTH = 8000  # 每个附件文本截断上限
MAX_ATTACHMENTS_PER_RECORD = 5


def _extract_pdf(path: Path) -> str:
    """用 PyMuPDF 提取 PDF 文本。"""
    if fitz is None:
        return ""
    text_parts: list[str] = []
    with fitz.open(str(path)) as pdf:
        for page in pdf:
            text_parts.append(page.get_text())
    return "\n".join(text_parts).strip()


def _extract_docx(path: Path) -> str:
    """用 python-docx 提取 Word 文本。"""
    if docx is None:
        return ""
    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs if p.text.strip()).strip()


def _extract_xlsx(path: Path) -> str:
    """用 openpyxl 提取 Excel 文本（所有 sheet 的所有单元格）。"""
    if openpyxl is None:
        return ""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    text_parts: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                text_parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(text_parts).strip()


def _extract_by_extension(path: Path) -> str:
    """根据扩展名选择提取器。"""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext in (".xlsx", ".xls"):
        return _extract_xlsx(path) if ext == ".xlsx" else ""
    # .doc 老格式不支持，跳过
    return ""


def _url_hash(url: str) -> str:
    return hashlib.sha256(url.encode()).hexdigest()


def _infer_extension(url: str) -> str:
    ext_match = re.search(r"\.(\w+)(?:[?#]|$)", url)
    return "." + ext_match.group(1).lower() if ext_match else ".bin"


def download_and_extract(url: str) -> tuple[str, str] | None:
    """下载附件到临时文件，提取文本，删除临时文件。

    返回 (file_type, extracted_text) 或 None（下载/提取失败）。
    不保留任何二进制文件。
    """
    if httpx is None:
        return None
    ext = _infer_extension(url)
    file_type = ext.lstrip(".")
    # 下载到临时文件
    try:
        with httpx.Client(follow_redirects=True, timeout=TIMEOUT_SECONDS) as client:
            resp = client.get(url, headers={"User-Agent": "BeaconOSINT/0.1"})
            if resp.status_code >= 400:
                return (file_type, "")
            if len(resp.content) > MAX_FILE_SIZE:
                return (file_type, "")
            # 写到临时文件，提取后自动删除
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(resp.content)
                tmp_path = Path(tmp.name)
            try:
                extracted = _extract_by_extension(tmp_path)
                return (file_type, extracted[:MAX_TEXT_LENGTH])
            finally:
                tmp_path.unlink(missing_ok=True)  # 立即删除临时文件
    except Exception:
        return None


def process_record_attachments(record: dict[str, Any]) -> dict[str, Any]:
    """处理单条记录的附件：下载 + 提取文本 + 删除文件。

    返回更新后的 record（加了 attachment_texts 字段）。
    不在磁盘上保留任何二进制文件。
    """
    links = record.get("attachment_links") or []
    if not links:
        record["attachment_texts"] = []
        record["attachment_count"] = 0
        return record

    texts: list[dict[str, str]] = []
    for url in links[:MAX_ATTACHMENTS_PER_RECORD]:
        result = download_and_extract(url)
        if result is None:
            continue
        file_type, extracted = result
        texts.append({
            "url": url,
            "url_hash": _url_hash(url),
            "file_type": file_type,
            "text": extracted,
            "text_length": len(extracted),
        })
    record["attachment_texts"] = texts
    record["attachment_count"] = len([t for t in texts if t["text"]])
    # 把附件文本追加到 content，让摘要/对比/印证能用上
    meaningful_texts = [t["text"] for t in texts if t["text"]]
    if meaningful_texts:
        existing_content = str(record.get("content") or "")
        attachment_text = "\n\n[附件内容]\n" + "\n\n".join(meaningful_texts)
        record["content"] = existing_content + attachment_text
    return record


def process_attachments(records: list[dict[str, Any]]) -> dict[str, int]:
    """批量处理附件。在 main.py 的 upsert_announcements 之前调用。

    下载到临时目录，提取文本后立即删除文件，不占用磁盘空间。
    """
    result = {
        "records_with_attachments": 0,
        "attachments_processed": 0,
        "texts_extracted": 0,
        "texts_empty": 0,
    }
    for record in records:
        links = record.get("attachment_links") or []
        if not links:
            continue
        result["records_with_attachments"] += 1
        process_record_attachments(record)
        for t in record.get("attachment_texts", []):
            result["attachments_processed"] += 1
            if t["text"]:
                result["texts_extracted"] += 1
            else:
                result["texts_empty"] += 1
    return result
